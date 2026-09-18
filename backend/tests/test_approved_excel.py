from io import BytesIO
from pathlib import Path

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from app.services.approved_excel import CELLS, parse_approved_excel

TEMPLATE = Path(__file__).parents[1] / "app/templates/approved-booklet.xlsx"


def workbook(changes=None, second=None):
    book = load_workbook(TEMPLATE)
    values = {
        "D3": "29/07/2026",
        "E3": "الساعه : 5:30 مساءا",
        "D5": "وصف عقار تجريبي",
        "D7": "فيلا",
        "F7": "٨٧٥٫٢٥",
        "D8": "001234",
        "D9": "005",
        "D10": "007",
        "F9": "الخطة",
        "F10": "****",
        "I8": "شارع",
        "I14": "٣٥٫٥",
        "C13": "معلومات إضافية",
    }
    for cell, value in {**values, **(changes or {})}.items():
        book.worksheets[0][cell] = value
    if second is not None:
        for cell, value in {**values, **second}.items():
            book.worksheets[1][cell] = value
    out = BytesIO()
    book.save(out)
    return out.getvalue()


def test_approved_cells_dates_and_empty_sheets():
    rows = parse_approved_excel(workbook(), "book.xlsx")
    assert len(rows) == 1
    assert not rows[0]["errors"]
    p = rows[0]["data"]["property_data"]
    assert p["area"] == "875.25"
    assert p["deed_number"] == "001234" and p["plot_number"] == "007"
    assert (
        p["auction_close_time"] == "17:30:00"
        and p["auction_close_date"] == "2026-07-29"
    )
    assert p["boundaries"]["north_length"] == "35.5"
    assert p["participation_amount"] is None and p["city"] == ""
    assert rows[0]["warnings"]


@pytest.mark.parametrize(
    "cell,value,code",
    [
        ("D5", None, "required"),
        ("D7", "", "required"),
        ("F7", -1, "number"),
        ("F7", 0, "number"),
        ("F7", "NaN", "number"),
        ("F10", "wrong", "number"),
        ("D3", "31/02/2026", "date"),
        ("E3", "25:99", "time"),
        ("I20", "javascript:alert(1)", "link"),
        ("D8", "=1+1", "formula"),
        ("D8", 1234567890123456, "identifier_precision"),
        ("C7", "wrong", "layout"),
    ],
)
def test_cell_errors(cell, value, code):
    row = parse_approved_excel(workbook({cell: value}), "a.xlsx")[0]
    assert any(i["cell"] == cell and i["code"] == code for i in row["errors"])


def test_duplicate_and_date_warnings():
    rows = parse_approved_excel(workbook(second={"D3": "29/07/2027"}), "a.xlsx")
    assert any(e["code"] == "duplicate" for e in rows[1]["errors"])
    assert all(any(w["code"] == "dates_differ" for w in r["warnings"]) for r in rows)


def test_hyperlinks_and_native_datetime():
    from datetime import datetime, time

    book = load_workbook(BytesIO(workbook()))
    sheet = book.worksheets[0]
    sheet["H20"].hyperlink = "https://example.com/survey"
    sheet["D3"] = datetime(2026, 7, 29)
    sheet["E3"] = time(17, 30)
    out = BytesIO()
    book.save(out)
    row = parse_approved_excel(out.getvalue(), "a.xlsx")[0]
    assert not row["errors"]
    assert row["data"]["property_data"]["survey_link"] == "https://example.com/survey"


@pytest.mark.parametrize(
    "data,name",
    [(b"wrong", "a.xlsx"), (b"pdf", "a.pdf"), (TEMPLATE.read_bytes(), "a.xls")],
)
def test_invalid_files(data, name):
    with pytest.raises(HTTPException) as exc:
        parse_approved_excel(data, name)
    assert exc.value.status_code == 400


def test_empty_template_and_limits():
    with pytest.raises(HTTPException):
        parse_approved_excel(TEMPLATE.read_bytes(), "a.xlsx")
    book = load_workbook(TEMPLATE)
    book.active["AE1"] = "too wide"
    out = BytesIO()
    book.save(out)
    with pytest.raises(HTTPException):
        parse_approved_excel(out.getvalue(), "a.xlsx")
    for sheet in load_workbook(TEMPLATE):
        assert sheet.sheet_view.rightToLeft
        assert all(sheet[c].value is None for c in CELLS)


def test_approved_api_atomic_commit_duplicates_and_download(admin):
    project = admin.post(
        "/api/projects", json={"name": "Approved template", "code": "APPROVED-1"}
    ).json()["id"]
    url = f"/api/projects/{project}/imports/approved/preview"
    download = admin.get("/api/booklet-import-template")
    assert download.status_code == 200 and download.content == TEMPLATE.read_bytes()
    bad = admin.post(
        url, files={"file": ("a.xlsx", workbook(second={"D8": "002", "F7": "bad"}))}
    ).json()
    assert bad["invalid_count"] == 1 and bad["valid_count"] == 1
    assert (
        admin.post(f"/api/projects/{project}/imports/{bad['id']}/commit").status_code
        == 400
    )
    assert not admin.get(f"/api/projects/{project}").json()["items"]
    data = workbook()
    previews = [
        admin.post(url, files={"file": ("a.xlsx", data)}).json() for _ in range(2)
    ]
    commit = f"/api/projects/{project}/imports/{previews[0]['id']}/commit"
    assert admin.post(commit).json() == {"imported": 1, "skipped": 0}
    assert admin.post(commit).status_code == 409
    assert (
        admin.post(
            f"/api/projects/{project}/imports/{previews[1]['id']}/commit"
        ).status_code
        == 409
    )
    repeated = admin.post(url, files={"file": ("a.xlsx", data)}).json()
    assert repeated["rows"][0]["errors"][0]["code"] == "already_imported"
    assert len(admin.get(f"/api/projects/{project}").json()["items"]) == 1


def test_common_city_and_extra_cells():
    rows = parse_approved_excel(workbook(), "a.xlsx", city="حائل")
    assert rows[0]["data"]["property_data"]["city"] == "حائل"
    assert not any(w["code"] == "city_missing" for w in rows[0]["warnings"])
    rows = parse_approved_excel(workbook({"J7": "unexpected data"}), "a.xlsx")
    assert any(e["code"] == "unsupported_cell" for e in rows[0]["errors"])


def test_approved_template_requires_login(client):
    assert client.get("/api/booklet-import-template").status_code == 401
