"""The approved booklet workbook: one property per worksheet, fixed input cells."""

import re
import zipfile
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from urllib.parse import urlsplit

from fastapi import HTTPException
from openpyxl import load_workbook
from pydantic import ValidationError

from .ingestion import normalize

LABELS = {
    "C3": "تغلق المزايده على العقار",
    "C5": "وصف العقار",
    "C7": "النوع",
    "E7": "المساحه م2",
    "C8": "رقم الصك",
    "E8": "الاستخدام",
    "C9": "رقم المخطط",
    "E9": "الحى",
    "C10": "رقم القطعة",
    "E10": "شيك الدخول",
    "C12": "معلومات اضافيه",
    "G8": "الحدود",
    "G14": "الاطوال",
    "H8": "شمال",
    "H9": "جنوب",
    "H10": "شرق",
    "H11": "غرب",
    "H14": "طول الحد الشمالي",
    "H15": "طول الحد الجنوبي",
    "H16": "طول الحد الشرقي",
    "H17": "طول الحد الغربي",
    "H20": "الرفع المساحى",
    "H21": "صور اضافيه",
    "H22": "اضغط هنا للوصول للرابط",
}
CELLS = {
    "D3": "auction_close_date",
    "E3": "auction_close_time",
    "D5": "description",
    "D7": "property_type",
    "F7": "area",
    "D8": "deed_number",
    "F8": "usage",
    "D9": "plan_number",
    "F9": "district",
    "D10": "plot_number",
    "F10": "participation_amount",
    "C13": "additional_information",
    "I8": "boundaries.north_description",
    "I9": "boundaries.south_description",
    "I10": "boundaries.east_description",
    "I11": "boundaries.west_description",
    "I14": "boundaries.north_length",
    "I15": "boundaries.south_length",
    "I16": "boundaries.east_length",
    "I17": "boundaries.west_length",
    "I20": "survey_link",
    "I21": "additional_images_link",
    "I22": "location_link",
}
REQUIRED = {"D5", "D7", "F7", "D8"}
DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹٫٬", "01234567890123456789.,")


def label(value):
    return re.sub(r"[\s\u064b-\u065fـ:]+", "", str(value or "")).translate(
        str.maketrans("أإآىة", "ااايه")
    )


def text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def missing(value):
    return not value or bool(re.fullmatch(r"[*/—–\-\s]+", value))


def identity(data):
    p = data.get("property_data", {})
    return tuple(
        text(p.get(k)).translate(DIGITS)
        for k in ("deed_number", "plan_number", "plot_number")
    )


def issue(cell, code, **context):
    return {"cell": cell, "code": code, "context": context}


def read_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    value = text(value).translate(DIGITS)
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(value, fmt).date().isoformat()
        except ValueError:
            continue
    raise ValueError("date")


def read_time(value):
    if isinstance(value, (datetime, time)):
        return (value.time() if isinstance(value, datetime) else value).isoformat()
    value = text(value).translate(DIGITS)
    value = re.sub(r"^\s*الساع[هة]\s*:?\s*", "", value)
    match = re.fullmatch(
        r"(\d{1,2}):(\d{2})(?::(\d{2}))?\s*(صباح[اًا]*|مساء[اًا]*|ص|م|AM|PM)?",
        value,
        re.I,
    )
    if not match:
        raise ValueError("time")
    h, m, s = map(int, (match[1], match[2], match[3] or 0))
    suffix = match[4]
    if suffix:
        if not 1 <= h <= 12:
            raise ValueError("time")
        h %= 12
        if suffix.lower() == "pm" or suffix.startswith("م"):
            h += 12
    return time(h, m, s).isoformat()


def parse_approved_excel(data, filename, city=""):
    if Path(filename).suffix.lower() != ".xlsx":
        raise HTTPException(400, "Approved template requires .xlsx")
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            if (
                sum(x.file_size for x in archive.infolist()) > 50 * 1024 * 1024
                or len(archive.infolist()) > 3000
            ):
                raise HTTPException(400, "Workbook expands beyond the 50 MB limit")
        book = load_workbook(BytesIO(data), data_only=False, keep_links=False)
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(
            400,
            "Could not read workbook. Check its format and remove password protection.",
        )
    try:
        if len(book.worksheets) > 100:
            raise HTTPException(400, "Approved workbook exceeds limits")
        rows, seen = [], set()
        for sheet in book:
            if sheet.max_row > 200 or sheet.max_column > 30:
                raise HTTPException(400, "Approved workbook exceeds limits")
            errors, warnings = [], []
            raw = {cell: text(sheet[cell].value) for cell in CELLS}
            # Hyperlinks may be attached to the label rather than the input cell.
            for target, anchor in [("I20", "H20"), ("I21", "H21"), ("I22", "H22")]:
                links = [
                    sheet[c].hyperlink.target
                    for c in (target, anchor)
                    if sheet[c].hyperlink and sheet[c].hyperlink.target
                ]
                if len(set(links)) > 1:
                    errors.append(issue(target, "conflicting_links"))
                if links:
                    raw[target] = links[0]
            for sheet_row in sheet:
                for cell in sheet_row:
                    if (
                        cell.coordinate not in CELLS
                        and cell.coordinate not in LABELS
                        and cell.value is not None
                    ):
                        errors.append(issue(cell.coordinate, "unsupported_cell"))
            mismatches = [
                c
                for c, expected in LABELS.items()
                if label(sheet[c].value) != label(expected)
            ]
            # Label anchors with hyperlink captions can vary, but must retain their labels.
            if mismatches:
                errors.extend(
                    issue(c, "layout", expected=LABELS[c]) for c in mismatches
                )
            elif not errors and not any(not missing(v) for v in raw.values()):
                continue  # Unused, intact template sheets are safe to omit.
            prop, mapped = {"boundaries": {}, "city": city.strip()}, {}
            for cell, field in CELLS.items():
                value = raw[cell]
                if sheet[cell].data_type == "f" or sheet[cell].data_type == "e":
                    errors.append(issue(cell, "formula"))
                    continue
                if missing(value):
                    if cell in REQUIRED:
                        errors.append(issue(cell, "required"))
                    elif cell in ("F10", "F9"):
                        warnings.append(issue(cell, "missing_optional"))
                    continue
                try:
                    if field in ("area", "participation_amount") or field.endswith(
                        "_length"
                    ):
                        number = Decimal(value.translate(DIGITS).replace(",", ""))
                        if (
                            not number.is_finite()
                            or number < 0
                            or (field == "area" and number == 0)
                        ):
                            raise ValueError("number")
                        value = str(number)
                    elif field.endswith("_date"):
                        value = read_date(sheet[cell].value)
                    elif field.endswith("_time"):
                        value = read_time(sheet[cell].value)
                    elif field.endswith("_link"):
                        parsed = urlsplit(value)
                        if (
                            parsed.scheme not in ("https", "http")
                            or not parsed.hostname
                            or parsed.username
                            or any(c.isspace() for c in value)
                        ):
                            raise ValueError("link")
                    elif field in ("deed_number", "plot_number", "plan_number"):
                        value = value.translate(DIGITS)
                        if (
                            re.fullmatch(r"0+", sheet[cell].number_format)
                            and value.isdigit()
                        ):
                            value = value.zfill(len(sheet[cell].number_format))
                        if (
                            sheet[cell].data_type == "n"
                            and len(value.replace(".", "")) > 15
                        ):
                            errors.append(issue(cell, "identifier_precision"))
                    if field == "description":
                        mapped[field] = value
                    elif field.startswith("boundaries."):
                        prop["boundaries"][field.split(".")[1]] = value
                    else:
                        prop[field] = value
                except (ValueError, InvalidOperation, OverflowError):
                    kind = (
                        "date"
                        if field.endswith("_date")
                        else "time"
                        if field.endswith("_time")
                        else "link"
                        if field.endswith("_link")
                        else "number"
                    )
                    errors.append(issue(cell, kind))
            mapped["title"] = (
                " · ".join(
                    v for v in (prop.get("property_type"), prop.get("district")) if v
                )
                or sheet.title
            )
            mapped["property_data"] = prop
            # City is intentionally absent from the source, never inferred from the filename.
            if not city.strip():
                warnings.append(issue("", "city_missing"))
            try:
                normalized = normalize(mapped)
            except ValidationError as exc:
                normalized = mapped
                for e in exc.errors():
                    field = ".".join(str(k) for k in e["loc"]).removeprefix(
                        "property_data."
                    )
                    cell = next((c for c, f in CELLS.items() if f == field), "")
                    errors.append(issue(cell, "invalid", field=field))
            key = identity(normalized)
            if key[0]:
                if key in seen:
                    errors.append(issue("D8", "duplicate"))
                seen.add(key)
            rows.append(
                {
                    "sheet": sheet.title,
                    "row_number": len(rows) + 1,
                    "data": normalized,
                    "source": raw,
                    "columns": [],
                    "errors": errors,
                    "issues": [],
                    "warnings": warnings,
                }
            )
        if not rows:
            raise HTTPException(400, "Approved workbook contains no properties")
        dates = {
            r["data"].get("property_data", {}).get("auction_close_date") for r in rows
        } - {None, ""}
        if len(dates) > 1:
            for row in rows:
                row["warnings"].append(issue("D3", "dates_differ"))
        return rows
    finally:
        book.close()
